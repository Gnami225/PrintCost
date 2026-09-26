"""Exports Excel : fiche de prix de revient et historique des calculs.

La fiche contient des formules vivantes : chaque montant du détail vaut
``quantité × prix unitaire``, la synthèse par poste additionne le détail
(``SOMME.SI``), les frais en pourcentage recalculent leur assiette. Modifier un
prix dans l'onglet Détail met donc à jour le total, comme dans l'application.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import config
from .historique import Enregistrement
from .models import UNITE_ASSIETTE, UNITE_HEURE, CostLine, PricingResult, is_production_poste, poste_rank
from .utils import unit_plural

FONT = "Arial"
ACCENT = "C8236A"
GRIS = "5E6169"
FILET = "E7E7EB"
FOND = "F6F6F8"

F_TITRE = Font(name=FONT, size=16, bold=True, color="2B2D33")
F_MARQUE = Font(name=FONT, size=10, bold=True, color=ACCENT)
F_LABEL = Font(name=FONT, size=10, color=GRIS)
F_TEXTE = Font(name=FONT, size=10, color="2B2D33")
F_GRAS = Font(name=FONT, size=10, bold=True, color="2B2D33")
F_ENTETE = Font(name=FONT, size=10, bold=True, color="2B2D33")
F_SAISIE = Font(name=FONT, size=10, color="0000FF")  # valeur figée (convention : bleu)
F_TOTAL = Font(name=FONT, size=12, bold=True, color=ACCENT)
F_NOTE = Font(name=FONT, size=9, italic=True, color=GRIS)
FILL_ENTETE = PatternFill("solid", fgColor=FOND)
BORD_BAS = Border(bottom=Side(style="thin", color="C9CAD1"))
BORD_HAUT = Border(top=Side(style="thin", color="2B2D33"))

FMT_MONTANT = '#,##0 "FCFA";-#,##0 "FCFA";"-"'
FMT_PRIX = '#,##0.00;-#,##0.00;"-"'
FMT_QTE = '#,##0.####;-#,##0.####;"-"'
FMT_HEURE = '0.0000'
FMT_PCT = '0.0%'
FMT_PCT_PRIX = '0.00%'
FMT_DATE = 'dd/mm/yyyy hh:mm'

ORIGINES = {
    "nomenclature": "Nomenclature",
    "gamme": "Gamme (temps machine)",
    "estimation": "Estimation automatique",
    "saisie": "Saisie",
    "paramètre": "Paramètres",
    "frais": "Frais paramétrés",
}


@dataclass
class Fiche:
    """Contenu d'une fiche de prix de revient (calcul en cours ou enregistré)."""

    reference: str
    date: datetime
    article_id: str
    designation: str
    nature: str
    categorie: str
    quantite: float
    unite: str
    lignes: list[CostLine]
    total: float
    parametres_version: str
    alertes: list[str] = field(default_factory=list)
    prix_indicatifs: list[str] = field(default_factory=list)
    numero: int | None = None


def fiche_depuis_resultat(result: PricingResult, reference: str = "", numero: int | None = None,
                          date: datetime | None = None) -> Fiche:
    art = result.article
    return Fiche(
        reference=reference, date=date or datetime.now().replace(microsecond=0),
        article_id=art.id, designation=art.designation, nature=art.nature, categorie=art.categorie,
        quantite=result.nb_exemplaires, unite=result.unit.singular, lignes=list(result.lines),
        total=result.total, parametres_version=result.parametres_version,
        alertes=list(result.alertes), prix_indicatifs=list(result.prix_indicatifs), numero=numero,
    )


def fiche_depuis_enregistrement(rec: Enregistrement) -> Fiche:
    indicatifs = sorted({l.libelle for l in rec.lignes if not l.prix_valide and not l.prix_manquant and l.montant > 0})
    return Fiche(
        reference=rec.reference, date=rec.cree_le, article_id=rec.article_id, designation=rec.designation,
        nature=rec.nature, categorie=rec.categorie, quantite=rec.quantite, unite=rec.unite,
        lignes=list(rec.lignes), total=rec.total, parametres_version=rec.parametres_version,
        alertes=list(rec.alertes), prix_indicatifs=indicatifs, numero=rec.id,
    )


def nom_fichier_fiche(fiche: Fiche) -> str:
    base = fiche.reference or f"{fiche.article_id}_{fiche.designation}"
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in base).strip("_")[:60]
    return f"Prix_de_revient_{safe or fiche.article_id}_{fiche.date:%Y%m%d}.xlsx"


# --------------------------------------------------------------------------- #
# Fiche de prix de revient
# --------------------------------------------------------------------------- #
def fiche_excel(fiche: Fiche) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Fiche"
    detail = wb.create_sheet("Détail")

    total_detail, n_last = _write_detail(detail, fiche)
    _write_synthese(ws, fiche, n_last)

    for sheet in (ws, detail):
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _write_detail(ws: Worksheet, fiche: Fiche) -> tuple[str, int]:
    headers = ["Poste", "Code", "Élément", "Quantité", "Unité", "Prix unitaire (FCFA)",
               "Montant (FCFA)", "Calcul", "Origine", "Prix validé"]
    widths = [22, 14, 42, 14, 16, 18, 16, 90, 22, 11]
    for col, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font, cell.fill, cell.border = F_ENTETE, FILL_ENTETE, BORD_BAS
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    first = 2
    rows_by_index: dict[int, int] = {}
    for i, line in enumerate(fiche.lignes):
        rows_by_index[i] = first + i
    last = first + len(fiche.lignes) - 1

    for i, line in enumerate(fiche.lignes):
        r = rows_by_index[i]
        ws.cell(row=r, column=1, value=line.poste).font = F_TEXTE
        ws.cell(row=r, column=2, value=line.code).font = F_TEXTE
        ws.cell(row=r, column=3, value=line.libelle).font = F_TEXTE
        q = ws.cell(row=r, column=4)
        if line.est_pourcentage:
            refs = [f"G{rows_by_index[j]}" for j, other in enumerate(fiche.lignes)
                    if j != i and _dans_assiette(other, line)]
            q.value = f"=SUM({','.join(refs)})" if refs else 0
            q.font = F_TEXTE
            q.number_format = FMT_MONTANT
            unite = "FCFA (assiette)"
        else:
            q.value = line.quantite
            q.font = F_SAISIE
            q.number_format = FMT_HEURE if line.unite == UNITE_HEURE else FMT_QTE
            unite = unit_plural(line.quantite, line.unite)
        ws.cell(row=r, column=5, value=unite).font = F_TEXTE
        pu = ws.cell(row=r, column=6, value=line.prix_unitaire)
        pu.font = F_SAISIE
        pu.number_format = FMT_PCT_PRIX if line.est_pourcentage else FMT_PRIX
        m = ws.cell(row=r, column=7, value=f"=D{r}*F{r}")
        m.font, m.number_format = F_TEXTE, FMT_MONTANT
        calc = ws.cell(row=r, column=8, value=line.formule)
        calc.font, calc.alignment = F_NOTE, Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=9, value=ORIGINES.get(line.origine, line.origine)).font = F_TEXTE
        valide = "Non renseigné" if line.prix_manquant else ("Oui" if line.prix_valide else "À valider")
        ws.cell(row=r, column=10, value=valide).font = F_TEXTE
        for col in range(1, 11):
            ws.cell(row=r, column=col).alignment = Alignment(
                vertical="top", wrap_text=col in (3, 8))

    total_row = max(last, first - 1) + 1
    ws.cell(row=total_row, column=3, value="Prix de revient total").font = F_GRAS
    total = ws.cell(row=total_row, column=7,
                    value=f"=SUM(G{first}:G{last})" if fiche.lignes else 0)
    total.font, total.number_format = F_GRAS, FMT_MONTANT
    for col in range(1, 11):
        ws.cell(row=total_row, column=col).border = BORD_HAUT
    note = ws.cell(row=total_row + 2, column=1,
                   value="En bleu : valeurs issues du calcul (quantités de la nomenclature et de la gamme, prix des "
                         "paramètres) ; modifiables, les montants et totaux se recalculent.")
    note.font = F_NOTE
    return f"'Détail'!G{total_row}", last


def _dans_assiette(other: CostLine, frais: CostLine) -> bool:
    """Lignes comprises dans l'assiette d'un frais en pourcentage (même règle que le moteur)."""
    if other.origine == "frais":
        return False
    if frais.assiette == "hors_frais":
        return True
    return is_production_poste(other.poste)


def _write_synthese(ws: Worksheet, fiche: Fiche, last_detail_row: int) -> None:
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 50

    ws["A1"] = "DIGIPRINT"
    ws["A1"].font = F_MARQUE
    ws["A2"] = "Fiche de prix de revient"
    ws["A2"].font = F_TITRE
    ws["A3"] = "Coût de revient hors marge commerciale, en FCFA."
    ws["A3"].font = F_NOTE

    infos = [
        ("Référence", fiche.reference or "Sans référence"),
        ("N° d'enregistrement", fiche.numero if fiche.numero is not None else "Non enregistré"),
        ("Date du calcul", fiche.date),
        ("Article", fiche.article_id),
        ("Désignation", fiche.designation),
        ("Nature", fiche.nature),
        ("Catégorie", fiche.categorie),
        ("Quantité", fiche.quantite),
        ("Version des paramètres", fiche.parametres_version),
    ]
    row = 5
    qty_cell = ""
    for label, value in infos:
        ws.cell(row=row, column=1, value=label).font = F_LABEL
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = F_TEXTE
        cell.alignment = Alignment(horizontal="left", wrap_text=True)
        if label == "Date du calcul":
            cell.number_format = FMT_DATE
        if label == "Quantité":
            cell.number_format = FMT_QTE
            ws.cell(row=row, column=3, value=unit_plural(fiche.quantite, fiche.unite)).font = F_TEXTE
            qty_cell = f"B{row}"
        row += 1

    row += 1
    for col, title in enumerate(["Poste", "Montant", "Part"], start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font, cell.fill, cell.border = F_ENTETE, FILL_ENTETE, BORD_BAS
    row += 1
    postes = sorted({line.poste for line in fiche.lignes}, key=poste_rank)
    first = row
    last = row + len(postes) - 1
    total_row = last + 1
    end = max(last_detail_row, 2)
    for poste in postes:
        ws.cell(row=row, column=1, value=poste).font = F_TEXTE
        m = ws.cell(row=row, column=2,
                    value=f"=SUMIF('Détail'!$A$2:$A${end},A{row},'Détail'!$G$2:$G${end})")
        m.font, m.number_format = F_TEXTE, FMT_MONTANT
        p = ws.cell(row=row, column=3, value=f"=IF($B${total_row}=0,0,B{row}/$B${total_row})")
        p.font, p.number_format = F_TEXTE, FMT_PCT
        row += 1
    ws.cell(row=total_row, column=1, value="Prix de revient total").font = F_GRAS
    t = ws.cell(row=total_row, column=2, value=f"=SUM(B{first}:B{last})" if postes else 0)
    t.font, t.number_format = F_TOTAL, FMT_MONTANT
    s = ws.cell(row=total_row, column=3, value=f"=SUM(C{first}:C{last})" if postes else 0)
    s.font, s.number_format = F_GRAS, FMT_PCT
    for col in range(1, 4):
        ws.cell(row=total_row, column=col).border = BORD_HAUT
    unit_row = total_row + 1
    ws.cell(row=unit_row, column=1,
            value=f"Coût unitaire (par {fiche.unite})").font = F_LABEL
    u = ws.cell(row=unit_row, column=2, value=f"=IF({qty_cell}=0,0,B{total_row}/{qty_cell})")
    u.font, u.number_format = F_GRAS, '#,##0.00 "FCFA"'

    row = unit_row + 2
    notes: list[str] = []
    if fiche.prix_indicatifs:
        notes.append("Prix indicatifs non encore validés : " + ", ".join(fiche.prix_indicatifs) + ".")
    notes.extend(fiche.alertes)
    if notes:
        ws.cell(row=row, column=1, value="Points d'attention").font = F_GRAS
        row += 1
        for text in notes:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            cell = ws.cell(row=row, column=1, value=text)
            cell.font = F_NOTE
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = 15 * max(1, len(text) // 110 + 1)
            row += 1
    ws.cell(row=row + 1, column=1,
            value=f"Édité par l'application {config.APP_TITLE}. Détail ligne à ligne dans l'onglet « Détail ».").font = F_NOTE


# --------------------------------------------------------------------------- #
# Historique
# --------------------------------------------------------------------------- #
def historique_excel(records: list[Enregistrement]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Historique"
    postes = sorted({p for rec in records for p in rec.postes}, key=poste_rank)
    headers = (["N°", "Date", "Référence", "Article", "Désignation", "Nature", "Catégorie",
                "Quantité", "Unité", "Prix de revient (FCFA)", "Coût unitaire (FCFA)"]
               + [f"{p} (FCFA)" for p in postes] + ["Version des paramètres"])
    widths = [7, 17, 26, 11, 22, 48, 28, 11, 14, 18, 16] + [16] * len(postes) + [14]
    for col, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font, cell.fill, cell.border = F_ENTETE, FILL_ENTETE, BORD_BAS
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32
    for r, rec in enumerate(records, start=2):
        values = [rec.id, rec.cree_le, rec.reference, rec.article_id, rec.designation, rec.nature,
                  rec.categorie, rec.quantite, unit_plural(rec.quantite, rec.unite), rec.total]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = F_TEXTE
        ws.cell(row=r, column=2).number_format = FMT_DATE
        ws.cell(row=r, column=8).number_format = FMT_QTE
        ws.cell(row=r, column=10).number_format = FMT_MONTANT
        unit = ws.cell(row=r, column=11, value=f"=IF(H{r}=0,0,J{r}/H{r})")
        unit.font, unit.number_format = F_TEXTE, '#,##0.00'
        for k, poste in enumerate(postes):
            cell = ws.cell(row=r, column=12 + k, value=round(rec.postes.get(poste, 0.0), 2))
            cell.font, cell.number_format = F_TEXTE, FMT_MONTANT
        ws.cell(row=r, column=12 + len(postes), value=rec.parametres_version).font = F_TEXTE
    ws.freeze_panes = "D2"
    if records:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(records) + 1}"
    ws.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
