import copy
import json

import pytest
from openpyxl import load_workbook

from digiprint.export import fiche_depuis_enregistrement, fiche_depuis_resultat, fiche_excel, historique_excel
from digiprint.historique import HistoriqueStore
from digiprint.models import PricingRequest
from digiprint.moteur import calculer
from digiprint.parametres import ParametresConflit, ParametresError, ParametresStore
from digiprint.stockage import FichiersStore, normalize_url
from tests.conftest import BASE


def test_parametres_persistance_et_sauvegardes(db, catalogue):
    store = ParametresStore(db, max_backups=2)
    p = store.load(catalogue)
    assert store.load(catalogue).version == p.version
    data = copy.deepcopy(p.data)
    for h in (200, 210, 220):
        data["general"]["hauteur_levee"] = h
        store.save(data, version_attendue=store.version_courante())
    assert len(store.sauvegardes()) == 2
    assert store.load(catalogue).general["hauteur_levee"] == 220
    with pytest.raises(ParametresConflit):
        store.save(data, version_attendue="ancienne")
    data["intrants"][0]["prix_achat"] = -1
    with pytest.raises(ParametresError):
        store.save(data)
    restaure = store.restaurer(store.sauvegardes()[-1]["id"], catalogue)
    assert restaure.general["hauteur_levee"] == 200


def test_document_illisible(db, catalogue):
    store = ParametresStore(db)
    store.load(catalogue)
    from sqlalchemy import update
    from digiprint.stockage import t_parametres
    with db.engine.begin() as conn:
        conn.execute(update(t_parametres).values(donnees="{pas du json"))
    p = store.load(catalogue)
    assert "_alerte" in p.data and store.sauvegardes()


def test_historique(db, catalogue, params):
    h = HistoriqueStore(db)
    r = calculer(catalogue, params, PricingRequest("ART-0001", 500, livraison=True, transport="MOTO", distance_km=3))
    n = h.ajouter(r, "Devis Société Générale")
    h.ajouter(calculer(catalogue, params, PricingRequest("ART-0044", 1000)), "Salon")
    assert [e.id for e in h.lister("societe")] == [n]
    rec = h.obtenir(n)
    assert rec.requete.distance_km == 3 and rec.total == pytest.approx(r.total)
    assert sum(l.montant for l in rec.lignes) == pytest.approx(rec.total)
    assert h.supprimer([n]) == 1 and h.compter() == 1


def test_fichiers(db):
    f = FichiersStore(db, max_versions=2)
    raw = BASE.read_bytes()
    for i in range(3):
        f.enregistrer(raw, f"v{i}.xlsx")
    assert [v.nom for v in f.versions()] == ["v2.xlsx", "v1.xlsx"]
    assert f.dernier().contenu == raw
    f.supprimer_tout()
    assert f.dernier() is None


def test_export_excel(db, catalogue, params, tmp_path):
    r = calculer(catalogue, params, PricingRequest("ART-0044", 1000, emballage="CARTON_M", emballage_nombre=2))
    wb = load_workbook(__import__("io").BytesIO(fiche_excel(fiche_depuis_resultat(r, "Test"))))
    detail = wb["Détail"]
    assert detail["G2"].value == "=D2*F2"
    assert any(str(c.value).startswith("=SUMIF") for row in wb["Fiche"].iter_rows() for c in row)
    h = HistoriqueStore(db)
    rec = h.obtenir(h.ajouter(r, "Test"))
    assert fiche_excel(fiche_depuis_enregistrement(rec))
    assert historique_excel(h.lister())


def test_adresse_neon():
    url = normalize_url("postgresql://u:p@ep-x-pooler.eu-central-1.aws.neon.tech/neondb")
    assert url.startswith("postgresql+psycopg2://") and "sslmode=require" in url
